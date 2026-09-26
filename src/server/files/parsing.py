import io

import pdfplumber
from docx import Document
from pptx import Presentation

PDF_FILE_TYPE = "pdf"
DOCX_FILE_TYPE = "docx"
PPTX_FILE_TYPE = "pptx"

SUPPORTED_FILE_TYPES = frozenset({PDF_FILE_TYPE, DOCX_FILE_TYPE, PPTX_FILE_TYPE})


class UnsupportedFileTypeError(Exception):
    pass


def parse_pdf(content: bytes) -> str:
    with pdfplumber.open(io.BytesIO(content)) as pdf:
        return "\n".join(page.extract_text() or "" for page in pdf.pages)


def parse_docx(content: bytes) -> str:
    document = Document(io.BytesIO(content))
    return "\n".join(paragraph.text for paragraph in document.paragraphs)


def parse_pptx(content: bytes) -> str:
    presentation = Presentation(io.BytesIO(content))
    return "\n".join(
        shape.text_frame.text
        for slide in presentation.slides
        for shape in slide.shapes
        if shape.has_text_frame
    )


PARSERS = {
    PDF_FILE_TYPE: parse_pdf,
    DOCX_FILE_TYPE: parse_docx,
    PPTX_FILE_TYPE: parse_pptx,
}


def file_type_from_filename(filename: str) -> str:
    suffix = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if suffix not in SUPPORTED_FILE_TYPES:
        raise UnsupportedFileTypeError(filename)
    return suffix


def extract_text(content: bytes, file_type: str) -> str:
    return PARSERS[file_type](content)
